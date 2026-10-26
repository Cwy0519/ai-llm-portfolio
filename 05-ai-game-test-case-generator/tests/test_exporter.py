# -*- coding: utf-8 -*-
"""导出层测试：Excel 结构、样式要素、CSV 编码、空用例保护。"""
from __future__ import annotations

import csv

import pytest
from openpyxl import load_workbook

from utils.case_model import HEADERS, build_cases
from utils.exporter import export_csv, export_xlsx


def sample_cases(n: int = 3):
    return build_cases(
        [
            {
                "title": f"用例{i}",
                "priority": ["P0", "P1", "P2"][i % 3],
                "type": "功能",
                "method": "场景法",
                "preconditions": "已登录",
                "steps": ["步骤一", "步骤二"],
                "expected": f"结果{i}",
            }
            for i in range(n)
        ],
        module="充值付费",
    )


def test_export_xlsx_creates_file(tmp_path):
    target = export_xlsx(sample_cases(), tmp_path / "cases.xlsx")
    assert target.exists() and target.stat().st_size > 0


def test_export_xlsx_headers_and_rows(tmp_path):
    cases = sample_cases(5)
    path = export_xlsx(cases, tmp_path / "cases.xlsx")
    sheet = load_workbook(path).active
    assert tuple(cell.value for cell in sheet[1]) == HEADERS
    assert sheet.max_row == len(cases) + 1
    assert sheet["A2"].value == "TC-001"


def test_export_xlsx_freezes_header_and_filters(tmp_path):
    path = export_xlsx(sample_cases(), tmp_path / "cases.xlsx")
    sheet = load_workbook(path).active
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == f"A1:I4"


def test_export_xlsx_has_summary_sheet(tmp_path):
    path = export_xlsx(sample_cases(), tmp_path / "cases.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["测试用例", "统计"]
    summary = workbook["统计"]
    assert summary["A1"].value == "指标"
    assert summary["A2"].value == "用例总数"


def test_export_csv_uses_utf8_sig(tmp_path):
    path = export_csv(sample_cases(2), tmp_path / "cases.csv")
    raw = path.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")  # BOM，Excel 打开不乱码
    with path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle))
    assert rows[0] == list(HEADERS)
    assert len(rows) == 3


def test_export_empty_raises(tmp_path):
    with pytest.raises(ValueError):
        export_xlsx([], tmp_path / "empty.xlsx")
    with pytest.raises(ValueError):
        export_csv([], tmp_path / "empty.csv")
