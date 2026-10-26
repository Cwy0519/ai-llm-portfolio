# -*- coding: utf-8 -*-
"""导出：把用例写成测试团队可以直接用的 Excel（含统计页）或 CSV。"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from utils.case_model import HEADERS, PRIORITIES, TestCase, stats

NAVY = "1F3F66"
LIGHT = "EEF3F9"
BORDER = Border(*[Side(style="thin", color="D9D9D9")] * 4)
COLUMN_WIDTHS = (11, 14, 36, 8, 10, 13, 30, 44, 38)


def export_xlsx(cases: Sequence[TestCase], path: str | Path, sheet_name: str = "测试用例") -> Path:
    """导出为标准测试用例表：首行冻结、带筛选、自动换行，另附统计页。"""
    if not cases:
        raise ValueError("没有用例可以导出")

    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name

    for index, width in enumerate(COLUMN_WIDTHS, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for col, title in enumerate(HEADERS, 1):
        cell = sheet.cell(row=1, column=col, value=title)
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    sheet.row_dimensions[1].height = 24

    for row_index, case in enumerate(cases, 2):
        for col, value in enumerate(case.row(), 1):
            cell = sheet.cell(row=row_index, column=col, value=value)
            cell.font = Font(name="宋体", size=10)
            cell.border = BORDER
            cell.alignment = Alignment(
                horizontal="center" if col in (1, 4, 5, 6) else "left",
                vertical="center",
                wrap_text=True,
            )
            if row_index % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=LIGHT)

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{len(cases) + 1}"

    # ---- 统计页 ----
    summary = workbook.create_sheet("统计")
    summary.column_dimensions["A"].width = 22
    summary.column_dimensions["B"].width = 12
    s = stats(cases)
    rows = [("用例总数", s["total"])]
    rows += [(f"优先级 {p}", s["by_priority"].get(p, 0)) for p in PRIORITIES]
    rows += [(f"类型 {t}", n) for t, n in s["by_type"].items() if n]
    rows += [(f"方法 {m}", n) for m, n in s["by_method"].items() if n]
    rows.append(("P0 占比", f"{s['p0_ratio'] * 100:.0f}%"))

    summary.append(["指标", "数值"])
    for cell in summary[1]:
        cell.font = Font(name="微软雅黑", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        summary.append(list(row))

    workbook.save(target)
    return target


def export_csv(cases: Sequence[TestCase], path: str | Path) -> Path:
    """导出 CSV（utf-8-sig，Excel 直接打开不乱码）。"""
    if not cases:
        raise ValueError("没有用例可以导出")
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    with target.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(HEADERS)
        for case in cases:
            writer.writerow(case.row())
    return target
